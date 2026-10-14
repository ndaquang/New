"""Xuất danh sách ra Excel (kèm ảnh thẻ) và vCard (nhập vào Danh bạ / Outlook / Google Contacts)."""

import io
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

XLSX_COLUMNS = [
    ("STT", None, 5), ("Ảnh", None, 26), ("Trạng thái", "status", 13), ("Phân loại", "category", 22), ("Họ tên", "name", 22),
    ("Chức vụ", "title", 26), ("Công ty", "company", 30), ("Di động", "mobile", 14), ("Điện thoại", "phone", 14),
    ("Fax", "fax", 14), ("Email", "email", 28), ("Website", "website", 22), ("Địa chỉ", "address", 40),
    ("Địa chỉ 2", "address2", 40), ("Ghi chú", "note", 24), ("Cần kiểm tra", "review", 16),
    ("File ảnh gốc", "source_file", 22), ("Ngày quét", "created_at", 17),
    ("Chữ nhận dạng (máy đọc)", "raw_text", 50),
]
HEADER_FILL = PatternFill("solid", fgColor="0B6BCB")
THIN = Side(style="thin", color="D0D7DE")


def _thumb(path: Path, width_px: int = 180) -> io.BytesIO | None:
    if not path.exists():
        return None
    im = Image.open(path).convert("RGB")
    im.thumbnail((width_px, width_px))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=80)
    buf.seek(0)
    return buf


def export_xlsx(cards: list[dict], images_dir: Path, out_path: Path, with_images: bool = True) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách"
    for col, (label, _, width) in enumerate(XLSX_COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width if (with_images or label != "Ảnh") else 4
    ws.row_dimensions[1].height = 24

    for i, card in enumerate(cards, 1):
        r = i + 1
        for col, (label, key, _) in enumerate(XLSX_COLUMNS, 1):
            if label == "STT":
                val = i
            elif key == "status":
                val = "Đã xác nhận" if card.get("status") == "confirmed" else "Chưa xác nhận"
            elif key == "created_at":
                val = datetime.fromtimestamp(card.get("created_at") or 0).strftime("%d/%m/%Y %H:%M")
            else:
                val = card.get(key, "") if key else ""
            c = ws.cell(row=r, column=col, value=val)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if key in ("mobile", "phone", "fax"):
                c.number_format = "@"  # giữ số 0 đầu
        thumb = _thumb(images_dir / (card.get("image") or "")) if with_images else None
        if thumb:
            img = XLImage(thumb)
            ws.add_image(img, f"B{r}")
            ws.row_dimensions[r].height = max(48, img.height * 0.75 + 6)
        else:
            ws.row_dimensions[r].height = 48

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(XLSX_COLUMNS))}{len(cards) + 1}"

    # Trang tổng hợp theo nhóm phân loại
    sm = wb.create_sheet("Tổng hợp")
    sm.append(["Phân loại", "Số danh thiếp"])
    for c in sm[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    for cat, n in Counter(c.get("category") or "Khác" for c in cards).most_common():
        sm.append([cat, n])
    sm.append(["Tổng cộng", len(cards)])
    sm[sm.max_row][0].font = sm[sm.max_row][1].font = Font(bold=True)
    sm.column_dimensions["A"].width = 36
    sm.column_dimensions["B"].width = 16

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _vesc(s: str) -> str:
    return str(s or "").replace("\\", "\\\\").replace("\n", "\\n").replace(",", "\\,").replace(";", "\\;")


def to_vcard(c: dict) -> str:
    words = (c.get("name") or "").split()
    family, given = (words[0], " ".join(words[1:])) if len(words) > 1 else ("", " ".join(words))
    out = ["BEGIN:VCARD", "VERSION:3.0", f"N:{_vesc(family)};{_vesc(given)};;;",
           f"FN:{_vesc(c.get('name') or c.get('company') or 'Danh thiếp')}"]
    if c.get("company"):
        out.append(f"ORG:{_vesc(c['company'])}")
    if c.get("title"):
        out.append(f"TITLE:{_vesc(c['title'])}")
    if c.get("mobile"):
        out.append(f"TEL;TYPE=CELL,VOICE:{c['mobile']}")
    if c.get("phone"):
        out.append(f"TEL;TYPE=WORK,VOICE:{c['phone']}")
    if c.get("fax"):
        out.append(f"TEL;TYPE=WORK,FAX:{c['fax']}")
    if c.get("email"):
        out.append(f"EMAIL;TYPE=INTERNET,WORK:{c['email']}")
    if c.get("website"):
        w = c["website"]
        out.append(f"URL:{w if w.lower().startswith('http') else 'http://' + w}")
    for key in ("address", "address2"):
        if c.get(key):
            out.append(f"ADR;TYPE=WORK:;;{_vesc(c[key])};;;;")
    note = "; ".join(x for x in (c.get("category"), c.get("note")) if x)
    if note:
        out.append(f"NOTE:{_vesc(note)}")
    out.append("END:VCARD")
    return "\r\n".join(out)


def export_vcf(cards: list[dict], out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\r\n".join(to_vcard(c) for c in cards) + "\r\n", encoding="utf-8")
    return out_path
