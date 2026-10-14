"""Xác nhận thẻ, bỏ thẻ trùng, chạy lại bỏ qua ảnh đã xử lý, xuất kho dữ liệu, chống treo."""

import csv
import json
import os

import pytest
from fastapi.testclient import TestClient

from app import server
from app.samples import SAMPLES, encode

os.environ["DANHTHIEP_NO_OPEN"] = "1"


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    server.init(tmp_path_factory.mktemp("wf"), use_vietocr=False, background=False)
    return TestClient(server.app)


def upload(client, name, fp="", skip_dup=""):
    data = encode(name, SAMPLES[name]())
    return client.post("/api/process", files={"file": (name, data, "image/jpeg")},
                       data={"fingerprint": fp, "skip_duplicates": skip_dup}).json()


def test_resume_skips_processed_files(client):
    upload(client, "anh_chup_ban_go.jpg", fp="a.jpg|1|1")
    done = client.post("/api/files/check", json={"fingerprints": ["a.jpg|1|1", "b.jpg|2|2"]}).json()["done"]
    assert done == ["a.jpg|1|1"]
    client.post("/api/files/forget")
    assert client.post("/api/files/check", json={"fingerprints": ["a.jpg|1|1"]}).json()["done"] == []


def test_skip_duplicates_while_scanning(client):
    before = len(client.get("/api/cards").json())
    res = upload(client, "anh_chup_ban_go.jpg", skip_dup="1")
    assert res["cards"] == [] and len(res["merged"]) == 1
    assert len(client.get("/api/cards").json()) == before


def test_find_and_merge_duplicates(client):
    upload(client, "anh_chup_ban_go.jpg")  # thêm 1 bản trùng
    cards = client.get("/api/cards").json()
    keep = cards[-1]["id"]
    client.post("/api/cards/confirm", json={"ids": [keep], "confirmed": True})
    d = client.get("/api/duplicates").json()
    assert d["groups"] and d["extra_cards"] >= 1
    res = client.post("/api/duplicates/merge").json()
    assert res["removed"] == d["extra_cards"]
    left = client.get("/api/cards").json()
    assert [c["id"] for c in left if "viteccons" in c["email"]] == [keep]  # giữ thẻ đã xác nhận
    assert client.get("/api/duplicates").json()["groups"] == []


def test_confirm_and_dw_export(client):
    upload(client, "danh_thiep_dien_tu.png")
    cards = client.get("/api/cards").json()
    cesco = next(c for c in cards if "cesco" in c["email"])
    assert cesco["status"] == ""
    r = client.post("/api/cards/confirm", json={"ids": [cesco["id"]], "confirmed": True}).json()
    assert r[0]["status"] == "confirmed" and r[0]["confirmed_at"] > 0

    # Phân tích lại không đụng vào thẻ đã xác nhận
    from app import server as srv

    srv.state.store.set_fields_silently(cesco["id"], {"company": "GIỮ NGUYÊN"})
    client.post("/api/reanalyze")
    assert next(c for c in client.get("/api/cards").json() if c["id"] == cesco["id"])["company"] == "GIỮ NGUYÊN"

    res = client.post("/api/export", json={"format": "dw", "only_confirmed": True, "images": True}).json()
    out = res["path"]
    rows = list(csv.DictReader(open(os.path.join(out, "cards.csv"), encoding="utf-8")))
    assert {r["status"] for r in rows} == {"confirmed"}
    row = next(r for r in rows if r["card_id"] == cesco["id"])
    assert row["mobile_e164"] == "+84918196121" and row["email_domain"] == "cesco.com.vn"
    assert row["created_at"].endswith("Z") and row["confirmed_at"].endswith("Z")
    lines = open(os.path.join(out, "cards.jsonl"), encoding="utf-8").read().splitlines()
    assert len(lines) == len(rows) and json.loads(lines[0])["card_id"]
    schema = json.load(open(os.path.join(out, "schema.json"), encoding="utf-8"))
    assert schema["primary_key"] == "card_id" and len(schema["columns"]) == len(rows[0])
    assert "CREATE TABLE" in open(os.path.join(out, "create_table.sql"), encoding="utf-8").read()
    assert os.listdir(os.path.join(out, "images"))

    # Xuất bổ sung: không có gì thay đổi thì báo lỗi rõ ràng
    r = client.post("/api/export", json={"format": "dw", "only_confirmed": True, "incremental": True})
    assert r.status_code == 400
    client.patch(f"/api/cards/{cesco['id']}", json={"fields": {"note": "gặp ở hội chợ"}})
    r = client.post("/api/export", json={"format": "dw", "only_confirmed": True, "incremental": True}).json()
    assert r["count"] == 1

    # Excel có cột Trạng thái
    from openpyxl import load_workbook

    x = client.post("/api/export", json={"format": "xlsx", "images": False}).json()
    ws = load_workbook(x["path"])["Danh sách"]
    assert ws["C1"].value == "Trạng thái" and "Đã xác nhận" in [c.value for c in ws["C"]]


def test_thumbnail(client):
    cid = client.get("/api/cards").json()[0]["id"]
    r = client.get(f"/api/cards/{cid}/thumb")
    assert r.status_code == 200 and len(r.content) < 60_000


def test_worker_recovers_from_hang():
    from app.worker import OCRWorker, WorkerTimeout

    w = OCRWorker(use_vietocr=False, timeout=3)
    w.start()
    try:
        assert w.run("sleep", 0.1) == "awake"
        with pytest.raises(WorkerTimeout):
            w.run("sleep", 30)  # giả lập ảnh làm treo
        assert w.restarts == 1
        name = "danh_thiep_dien_tu.png"
        cards = w.run("process", (encode(name, SAMPLES[name]()), server.categories()), timeout=120)
        assert len(cards) == 1 and cards[0]["email"] == "long.vu@cesco.com.vn"
    finally:
        w.stop()
