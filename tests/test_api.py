import os

import pytest
from fastapi.testclient import TestClient

from app import server
from app.samples import SAMPLES, encode

os.environ["DANHTHIEP_NO_OPEN"] = "1"


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    # Bộ đọc RapidOCR là đủ để kiểm tra luồng cắt thẻ / lưu / xuất file.
    server.init(tmp_path_factory.mktemp("data"), use_vietocr=False, background=False)
    return TestClient(server.app)


def test_status(client):
    s = client.get("/api/status").json()
    assert s["ready"] and "Nhà thầu xây dựng" in s["categories"]


def test_process_export_flow(client, tmp_path):
    name = "scan_hai_the.jpg"
    r = client.post("/api/process", files={"file": (name, encode(name, SAMPLES[name]()), "image/jpeg")}, data={"batch": "b1"})
    assert r.status_code == 200, r.text
    cards = r.json()["cards"]
    assert len(cards) == 2
    emails = {c["email"] for c in cards}
    assert emails == {"quang.nguyen@viteccons.vn", "long.vu@cesco.com.vn"}
    assert all(c["source_file"] == name for c in cards)

    # ảnh thẻ đã cắt
    img = client.get(f"/api/cards/{cards[0]['id']}/image")
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"

    # sửa tay
    cid = cards[0]["id"]
    upd = client.patch(f"/api/cards/{cid}", json={"fields": {"name": "Tên Đã Sửa", "id": "hack"}}).json()
    assert upd["name"] == "Tên Đã Sửa" and upd["id"] == cid

    # thẻ trùng được đánh dấu
    again = client.post("/api/process", files={"file": (name, encode(name, SAMPLES[name]()), "image/jpeg")}).json()["cards"]
    assert all("trùng" in c["review"] for c in again)

    # xuất Excel + vCard
    from openpyxl import load_workbook

    x = client.post("/api/export", json={"format": "xlsx"}).json()
    wb = load_workbook(x["path"])
    ws = wb["Danh sách"]
    assert ws.max_row == 5 and ws["D1"].value == "Họ tên"
    assert len(ws._images) == 4
    assert wb["Tổng hợp"]["B" + str(wb["Tổng hợp"].max_row)].value == 4
    v = client.post("/api/export", json={"format": "vcf", "ids": [cid]}).json()
    text = open(v["path"], encoding="utf-8").read()
    assert text.count("BEGIN:VCARD") == 1 and "FN:Tên Đã Sửa" in text

    # xoay thẻ, đổi phân loại, xóa
    assert client.post(f"/api/cards/{cid}/rotate").status_code == 200
    client.post("/api/cards/category", json={"ids": [cid], "category": "Khác"})
    assert next(c for c in client.get("/api/cards").json() if c["id"] == cid)["category"] == "Khác"
    client.post("/api/cards/delete", json={"ids": [c["id"] for c in again]})
    assert len(client.get("/api/cards").json()) == 2


def test_bad_image(client):
    r = client.post("/api/process", files={"file": ("x.jpg", b"not an image", "image/jpeg")})
    assert r.status_code == 400


def test_learning_and_calibration_pack(client):
    import json
    import zipfile

    name = "anh_chup_ban_go.jpg"
    first = client.post("/api/process", files={"file": (name, encode(name, SAMPLES[name]()), "image/jpeg")}).json()["cards"][0]
    assert first["website"] == "www.viteccons.vn"
    # Bạn sửa tên công ty và phân loại cho thẻ này
    client.patch(f"/api/cards/{first['id']}", json={"fields": {"company": "Công ty CP Viteccons", "category": "Tư vấn thiết kế / giám sát / QS"}})
    # Thẻ sau cùng tên miền tự dùng giá trị đã sửa
    second = client.post("/api/process", files={"file": (name, encode(name, SAMPLES[name]()), "image/jpeg")}).json()["cards"][0]
    assert second["company"] == "Công ty CP Viteccons"
    assert second["category"] == "Tư vấn thiết kế / giám sát / QS"

    res = client.post("/api/export", json={"format": "calib"}).json()
    with zipfile.ZipFile(res["path"]) as z:
        items = json.loads(z.read("the-da-sua.json"))
        mine = next(i for i in items if i["id"] == first["id"])
        assert set(mine["fields_changed"]) >= {"company", "category"}
        assert mine["auto"]["company"] == "VITECCONS" and mine["raw_text"]
        assert f"anh/{first['id']}.jpg" in z.namelist()


def test_reanalyze_keeps_manual_edits(client):
    name = "danh_thiep_dien_tu.png"
    card = client.post("/api/process", files={"file": (name, encode(name, SAMPLES[name]()), "image/png")}).json()["cards"][0]
    # Giả lập kết quả cũ, sai (như bản trước): email trống, công ty sai; bạn đã sửa tay chức vụ.
    from app import server as srv

    srv.state.store.set_fields_silently(card["id"], {"email": "", "company": "SAI"})
    import json as _json

    auto = _json.loads(card["auto_fields"])
    auto.update({"email": "", "company": "SAI"})
    srv.state.store.set_fields_silently(card["id"], {"auto_fields": _json.dumps(auto, ensure_ascii=False)})
    client.patch(f"/api/cards/{card['id']}", json={"fields": {"title": "Chức vụ tôi tự sửa"}})

    res = client.post("/api/reanalyze").json()
    assert res["cards"] >= 1
    after = next(c for c in client.get("/api/cards").json() if c["id"] == card["id"])
    assert after["email"] == card["email"] and after["company"] == card["company"]  # được sửa lại
    assert after["title"] == "Chức vụ tôi tự sửa"  # ô sửa tay giữ nguyên
